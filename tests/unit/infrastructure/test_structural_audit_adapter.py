from __future__ import annotations

import base64

from docx import Document
from docx.shared import Inches

from docs.application.structural_audit import StructuralAuditService
from docs.domain.review import ReviewDimension
from docs.infrastructure.audit.structural_audit_adapter import StructuralAuditAdapter


def test_docx_structural_audit_checks_declarative_structure(tmp_path):
    docx_path = tmp_path / "report.docx"
    document = Document()
    document.core_properties.title = "Harness report"
    document.add_heading("INTRODUCTION", level=1)
    document.add_heading("METHODS", level=1)
    document.add_table(rows=1, cols=1)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    document.save(docx_path)

    result = StructuralAuditService(StructuralAuditAdapter()).audit(
        docx_path,
        {
            "headings": ["INTRODUCTION", "METHODS"],
            "sections": ["INTRODUCTION", "METHODS"],
            "tables": {"minimum": 1},
            "images": {"minimum": 0},
            "captions": {"required": False},
            "references": {"required": False},
            "metadata": {"required": ["title"]},
            "page_size": (612, 792),
        },
    )

    assert result.passed is True
    assert result.issues == []


def test_docx_structural_audit_reports_missing_declared_requirements(tmp_path):
    docx_path = tmp_path / "report.docx"
    document = Document()
    document.add_heading("METHODS", level=1)
    document.save(docx_path)

    result = StructuralAuditService(StructuralAuditAdapter()).audit(
        docx_path,
        {
            "headings": ["INTRODUCTION"],
            "tables": {"minimum": 1},
            "references": {"required": True},
            "metadata": {"required": ["title"]},
        },
    )

    assert {issue.code for issue in result.issues} >= {
        "structure.headings.order",
        "structure.tables.minimum",
        "structure.references.missing",
        "structure.metadata.missing",
    }
    assert all(issue.dimension is ReviewDimension.STRUCTURAL for issue in result.issues)


def test_structural_audit_accepts_json_page_size_list(tmp_path):
    docx_path = tmp_path / "report.docx"
    document = Document()
    document.sections[0].page_width = Inches(8.5)
    document.sections[0].page_height = Inches(11)
    document.save(docx_path)

    result = StructuralAuditService(StructuralAuditAdapter()).audit(docx_path, {"page_size": [600, 790]})

    assert any(issue.code == "structure.page_size" for issue in result.issues)


def test_docx_structural_audit_classifies_missing_captions_as_accessibility(tmp_path):
    image_path = tmp_path / "image.png"
    image_path.write_bytes(
        base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL6OwAAAABJRU5ErkJggg==")
    )
    docx_path = tmp_path / "report.docx"
    document = Document()
    document.add_picture(str(image_path))
    document.save(docx_path)

    result = StructuralAuditService(StructuralAuditAdapter()).audit(
        docx_path, {"captions": {"required": True}}
    )

    assert [(issue.code, issue.dimension) for issue in result.issues] == [
        ("structure.captions.missing", ReviewDimension.ACCESSIBILITY)
    ]


def test_docx_structural_audit_executes_template_contract_geometry_and_components(tmp_path):
    docx_path = tmp_path / "report.docx"
    document = Document()
    document.add_heading("METHODS", level=1)
    section = document.sections[0]
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.5)
    document.save(docx_path)

    result = StructuralAuditService(StructuralAuditAdapter()).audit(docx_path, {
        "page_geometry": {
            "size": "letter",
            "orientation": "portrait",
            "margins_cm": {"top": 2.5, "right": 2.5, "bottom": 2.5, "left": 2.5},
        },
        "components": [
            {"kind": "headings", "items": ["INTRODUCTION"]},
            {"kind": "tables", "minimum": 1},
            {"kind": "references", "required": True},
            {"kind": "metadata", "required": ["title"]},
        ],
    })

    assert {issue.code for issue in result.issues} >= {
        "contract.page_geometry.orientation",
        "contract.page_geometry.margins",
        "structure.headings.order",
        "structure.tables.minimum",
        "structure.references.missing",
        "structure.metadata.missing",
    }


def test_docx_structural_audit_executes_assets_slots_fidelity_and_degradations(tmp_path):
    docx_path = tmp_path / "report.docx"
    document = Document()
    document.add_heading("INTRODUCTION", level=1)
    document.add_paragraph("[[slot:title]]")
    document.save(docx_path)

    result = StructuralAuditService(StructuralAuditAdapter()).audit(docx_path, {
        "required_assets": [{"id": "logo", "path": "missing-logo.png", "required": True}],
        "editable_slots": [{"id": "title", "required": True}],
        "fidelity_checks": [{"id": "page-count", "minimum": 2}],
        "allowed_degradations": ["required asset logo", "fidelity page-count"],
    })

    findings = {issue.code: issue for issue in result.issues}
    assert findings["contract.required_assets.logo"].severity == "warning"
    assert findings["contract.editable_slots.title"].severity == "info"
    assert findings["contract.fidelity_checks.page-count"].severity == "warning"
