from __future__ import annotations

import base64
import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches

import docs.application.docx_import as docx_import
from docs.application.docx_import import DocxImportService, ImportReport

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUB"
    "AScY42YAAAAASUVORK5CYII="
)


def _make_docx(path: Path) -> bytes:
    image = path.with_suffix(".png")
    image.write_bytes(_PNG)
    document = Document()
    document.core_properties.title = "Structured import fixture"
    document.core_properties.author = "Ada Lovelace"
    document.styles.add_style("Architecture Note", WD_STYLE_TYPE.PARAGRAPH)
    document.add_heading("System overview", level=1)
    paragraph = document.add_paragraph("The service preserves structure.")
    paragraph.style = "Architecture Note"
    document.add_picture(str(image), width=Inches(0.25))
    caption = document.add_paragraph("Figure 1. Import boundary")
    caption.style = "Caption"
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Importer"
    table.cell(1, 1).text = "Application"
    section = document.sections[0]
    section.header.paragraphs[0].text = "Confidential"
    section.footer.paragraphs[0].text = "Internal draft"
    document.add_paragraph("")
    document.save(path)
    return path.read_bytes()


def _tree_bytes(root: Path) -> dict[str, bytes]:
    return {
        path.relative_to(root).as_posix(): path.read_bytes()
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def test_import_preserves_source_and_extracts_structured_deterministic_outputs(tmp_path: Path) -> None:
    source = tmp_path / "input.docx"
    source_bytes = _make_docx(source)
    output = tmp_path / "imported"

    report = DocxImportService().import_file(source, output)

    assert report.source_hash == hashlib.sha256(source_bytes).hexdigest()
    assert (output / "source.docx").read_bytes() == source_bytes
    markdown = (output / "document.md").read_text(encoding="utf-8")
    assert "# System overview" in markdown
    assert "<!-- style: Architecture Note -->\nThe service preserves structure." in markdown
    assert "![Image 1](assets/" in markdown
    assert "*Figure 1. Import boundary*" in markdown
    assert "| Component | Owner |" in markdown
    assert "## Headers\n\n### Section 1\n\nConfidential" in markdown
    assert "## Footers\n\n### Section 1\n\nInternal draft" in markdown
    assert 'title: "Structured import fixture"' in markdown
    assert 'author: "Ada Lovelace"' in markdown
    assets = list((output / "assets").iterdir())
    assert len(assets) == 1
    assert assets[0].read_bytes() == _PNG
    assert "source:exact-bytes" in report.preserved
    assert "paragraph-style:Architecture Note" in report.normalized
    assert "empty-paragraph" in report.dropped

    payload = json.loads((output / "import-report.json").read_text(encoding="utf-8"))
    assert payload == {
        "dropped": list(report.dropped),
        "normalized": list(report.normalized),
        "preserved": list(report.preserved),
        "source_hash": report.source_hash,
        "unsupported": list(report.unsupported),
    }
    assert report == ImportReport(**{**payload, "preserved": tuple(payload["preserved"]),
                                     "normalized": tuple(payload["normalized"]),
                                     "dropped": tuple(payload["dropped"]),
                                     "unsupported": tuple(payload["unsupported"])})


def test_import_is_byte_deterministic_for_the_same_source(tmp_path: Path) -> None:
    source = tmp_path / "input.docx"
    _make_docx(source)
    first = tmp_path / "first"
    second = tmp_path / "second"

    first_report = DocxImportService().import_file(source, first)
    second_report = DocxImportService().import_file(source, second)

    assert first_report == second_report
    assert _tree_bytes(first) == _tree_bytes(second)


def test_import_reports_unsupported_ooxml_without_claiming_lossless_conversion(tmp_path: Path) -> None:
    source = tmp_path / "input.docx"
    _make_docx(source)
    with ZipFile(source, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/footnotes.xml", "<w:footnotes xmlns:w='urn:unsupported'/>")

    report = DocxImportService().import_file(source, tmp_path / "imported")

    assert "ooxml-part:word/footnotes.xml" in report.unsupported
    assert all("lossless" not in item.lower() for item in report.preserved)


def test_missing_python_docx_has_a_clear_failure_and_publishes_nothing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "input.docx"
    source.write_bytes(b"not inspected")
    output = tmp_path / "imported"

    def missing_dependency(name: str):
        if name == "docx":
            raise ModuleNotFoundError("No module named 'docx'")
        raise AssertionError(name)

    monkeypatch.setattr(docx_import, "import_module", missing_dependency)

    with pytest.raises(RuntimeError, match="python-docx is required"):
        DocxImportService().import_file(source, output)
    assert not output.exists()


def test_malformed_docx_preserves_existing_output_without_partial_publication(tmp_path: Path) -> None:
    source = tmp_path / "broken.docx"
    source.write_bytes(b"not a zip package")
    output = tmp_path / "imported"
    output.mkdir()
    sentinel = output / "sentinel.txt"
    sentinel.write_text("previous", encoding="utf-8")

    with pytest.raises(ValueError, match="Malformed DOCX"):
        DocxImportService().import_file(source, output)

    assert _tree_bytes(output) == {"sentinel.txt": b"previous"}
