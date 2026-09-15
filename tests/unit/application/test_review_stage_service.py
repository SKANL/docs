from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document

from docs.application.format_audit import FormatAuditService
from docs.application.review import ReviewService
from docs.application.review_stages import ReviewStageService
from docs.application.structural_audit import StructuralAuditService
from docs.domain.models.template import Template
from docs.domain.pipeline_policy import PipelineMode, PipelinePolicy
from docs.domain.review import Issue, ReviewDimension, ReviewResult
from docs.domain.workspace import Workspace
from docs.infrastructure.audit.structural_audit_adapter import StructuralAuditAdapter
from docs.infrastructure.docx.python_docx_audit_adapter import PythonDocxAuditAdapter
from docs.infrastructure.persistence.json_section_repository import JsonSectionRepository


def _service(tmp_path: Path) -> ReviewStageService:
    workspace = Workspace(documents_dir=tmp_path / "documents", templates_dir=tmp_path / "templates")
    return ReviewStageService(
        structural_audit=StructuralAuditService(StructuralAuditAdapter()),
        format_audit=FormatAuditService(PythonDocxAuditAdapter()),
        document_review=ReviewService(JsonSectionRepository(workspace)),
        rules_manifest_state=lambda _config: (True, 1),
    )


def _document(tmp_path: Path, heading: str = "INTRODUCCION") -> Path:
    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph("Contenido de prueba.")
    path = tmp_path / "fixture.docx"
    document.save(path)
    return path


def _template() -> Template:
    return Template(type="fixture", title="Fixture")


def test_review_stages_run_real_adapters_for_a_compliant_fixture_document(tmp_path: Path) -> None:
    artifact = _document(tmp_path)
    service = _service(tmp_path)

    outcomes = service.run_all(
        document_id="fixture",
        artifact_path=artifact,
        config={"apa7": {"citation_style": "none"}},
        template=_template(),
        policy=PipelinePolicy(PipelineMode.draft),
        rebuild=lambda output: shutil.copyfile(artifact, output),
        scratch_dir=tmp_path / "rebuild",
    )

    assert tuple(outcomes) == (
        "structural-audit",
        "editorial-review",
        "evidence-review",
        "consistency-review",
        "accessibility-review",
        "visual-review",
        "reproducibility-check",
    )
    assert all(outcome.ok for outcome in outcomes.values())
    assert all(not outcome.warnings for outcome in outcomes.values())


def test_blocking_visual_finding_degrades_only_in_draft_policy(tmp_path: Path) -> None:
    artifact = _document(tmp_path, heading="Introduccion")
    service = _service(tmp_path)
    config = {"apa7": {"citation_style": "none"}}

    draft = service.visual_review(artifact, config, PipelinePolicy(PipelineMode.draft))
    strict = service.visual_review(artifact, config, PipelinePolicy(PipelineMode.strict))
    release = service.visual_review(artifact, config, PipelinePolicy(PipelineMode.release))

    assert draft.ok is True
    assert draft.warnings == ("Título de primer orden no está en mayúsculas sostenidas: `Introduccion`.",)
    assert strict.ok is False
    assert set(draft.warnings).issubset(strict.errors)
    assert release.ok is False
    assert set(draft.warnings).issubset(release.errors)


def test_run_stage_executes_only_the_requested_review_adapter(tmp_path: Path) -> None:
    artifact = _document(tmp_path)
    service = _service(tmp_path)

    outcome = service.run_stage(
        "visual-review",
        document_id="fixture",
        artifact_path=artifact,
        config={"apa7": {"citation_style": "none"}},
        template=_template(),
        policy=PipelinePolicy(PipelineMode.draft),
        rebuild=lambda _output: (_ for _ in ()).throw(AssertionError("unexpected rebuild")),
        scratch_dir=tmp_path / "rebuild",
    )

    assert outcome.ok is True


def test_evidence_review_applies_pipeline_policy_to_evidence_findings(tmp_path: Path) -> None:
    class Review:
        def review_document(self, *args: object, **kwargs: object) -> ReviewResult:
            return ReviewResult([
                Issue(
                    "warning",
                    "evidence needs support",
                    code="evidence.missing_support",
                    dimension=ReviewDimension.EVIDENCE,
                )
            ])

    service = ReviewStageService(
        structural_audit=StructuralAuditService(StructuralAuditAdapter()),
        format_audit=FormatAuditService(PythonDocxAuditAdapter()),
        document_review=Review(),
        rules_manifest_state=lambda _config: (True, 1),
    )

    outcome = service.run_stage(
        "evidence-review",
        document_id="fixture",
        artifact_path=tmp_path / "unused.docx",
        config={},
        template=_template(),
        policy=PipelinePolicy(PipelineMode.strict),
        rebuild=lambda output: output,
        scratch_dir=tmp_path / "rebuild",
    )

    assert outcome.ok is False
    assert outcome.errors == ("evidence needs support",)
    assert not outcome.warnings


def _multiformat_service(tmp_path: Path, **options) -> ReviewStageService:
    from docs.application.render_verification import RenderVerificationService
    from docs.infrastructure.verification.render_verification_adapter import RenderVerificationAdapter

    workspace = Workspace(documents_dir=tmp_path / "documents", templates_dir=tmp_path / "templates")
    return ReviewStageService(
        structural_audit=StructuralAuditService(StructuralAuditAdapter()),
        format_audit=FormatAuditService(PythonDocxAuditAdapter()),
        document_review=ReviewService(JsonSectionRepository(workspace)),
        rules_manifest_state=lambda _: (True, 1),
        render_verification=RenderVerificationService(RenderVerificationAdapter()),
        **options,
    )


def test_html_accessibility_stage_uses_static_checks_not_docx_audit(tmp_path):
    artifact = tmp_path / "document.html"
    artifact.write_text('<html><body><h2>Content</h2></body></html>')
    outcome = _multiformat_service(tmp_path).accessibility_review(artifact, {}, PipelinePolicy())
    assert not outcome.ok
    assert "language" in outcome.detail and "h1" in outcome.detail and "landmark" in outcome.detail


def test_accessibility_and_visual_dimensions_are_independent(tmp_path):
    artifact = tmp_path / "document.html"
    artifact.write_text('<html lang="en"><body><header><h1>Title</h1></header>'
                        '<main><h2>Content</h2><p>Text.</p></main></body></html>')
    service = _multiformat_service(tmp_path)
    accessibility = service.accessibility_review(artifact, {}, PipelinePolicy(PipelineMode.strict))
    visual = service.visual_review(artifact, {}, PipelinePolicy(PipelineMode.strict))
    assert accessibility.ok and not accessibility.warnings
    assert not visual.ok and "Browser renderer unavailable" in visual.detail
    assert "WARNING" not in visual.detail  # Report and gate must reflect the same policy.


def test_release_allows_static_html_qa_when_browser_renderer_is_unavailable(tmp_path):
    artifact = tmp_path / "document.html"
    artifact.write_text('<html lang="en"><body><header><h1>Title</h1></header>'
                        '<main><p>Text.</p></main></body></html>')

    outcome = _multiformat_service(tmp_path).visual_review(
        artifact, {}, PipelinePolicy(PipelineMode.release)
    )

    assert outcome.ok
    assert any("Browser renderer unavailable" in warning for warning in outcome.warnings)


def test_pdf_tag_warning_degrades_only_according_to_policy(tmp_path):


    artifact = tmp_path / "document.pdf"
    _write_blank_pdf(artifact)
    service = _multiformat_service(tmp_path)
    draft = service.accessibility_review(artifact, {}, PipelinePolicy())
    assert draft.ok and any("tagged structure" in message for message in draft.warnings)
    for policy in (PipelinePolicy(PipelineMode.strict), PipelinePolicy(PipelineMode.release),
                   PipelinePolicy(warning_codes=("accessibility.pdf.untagged",))):
        assert not service.accessibility_review(artifact, {}, policy).ok
    artifact.write_bytes(b"not a pdf")
    assert not service.accessibility_review(artifact, {}, PipelinePolicy()).ok


def test_visual_stage_reuses_pdf_renderer_and_preview_directory(tmp_path):


    artifact = tmp_path / "document.pdf"
    _write_blank_pdf(artifact)
    previews = tmp_path / "qa"
    result = _multiformat_service(tmp_path).visual_review(
        artifact, {"paths": {"output_qa_dir": str(previews)}}, PipelinePolicy(),
    )
    assert not result.ok and "vacía" in result.detail
    assert list(previews.rglob("*.png"))


def test_visual_stage_honors_blank_page_profile(tmp_path):


    artifact = tmp_path / "document.pdf"
    _write_blank_pdf(artifact)
    config = {"visual_qa": {"allow_blank_pages": True, "expected_page_size": [612, 792]}}
    service = _multiformat_service(tmp_path)
    assert service.visual_review(artifact, config, PipelinePolicy()).ok
    assert not service.visual_review(artifact, config, PipelinePolicy(PipelineMode.strict)).ok
    config["visual_qa"]["expected_page_size"] = [100, 100]
    assert not service.visual_review(artifact, config, PipelinePolicy()).ok


def test_multiformat_stage_without_render_port_reports_capability_gap(tmp_path):
    artifact = tmp_path / "document.html"
    artifact.write_text('<html><body>text</body></html>')
    for operation in (_service(tmp_path).accessibility_review, _service(tmp_path).visual_review):
        draft = operation(artifact, {}, PipelinePolicy())
        strict = operation(artifact, {}, PipelinePolicy(PipelineMode.strict))
        assert draft.ok and "not configured" in draft.detail
        assert not strict.ok


def _write_blank_pdf(path):
    import pypdfium2 as pdfium

    with pdfium.PdfDocument.new() as document:
        page = document.new_page(612, 792)
        document.save(path)
        page.close()


def test_missing_artifact_is_not_degraded_as_missing_renderer(tmp_path):
    result = _service(tmp_path).visual_review(tmp_path / "missing.pdf", {}, PipelinePolicy())
    assert not result.ok


def test_visual_profile_does_not_treat_false_string_as_permission_for_blank_pages(tmp_path):
    artifact = tmp_path / "document.pdf"
    _write_blank_pdf(artifact)
    result = _multiformat_service(tmp_path).visual_review(
        artifact, {"visual_qa": {"allow_blank_pages": "false"}}, PipelinePolicy(),
    )
    assert not result.ok and "profile" in result.detail.lower()


def test_visual_stage_compares_configured_pdf_baseline_without_updating_it(tmp_path):
    from PIL import Image

    artifact = tmp_path / "document.pdf"
    _write_blank_pdf(artifact)
    baseline = tmp_path / "baseline"
    baseline.mkdir()
    baseline_page = baseline / "document-p01.png"
    Image.new("RGB", (1275, 1650), "black").save(baseline_page)
    original = baseline_page.read_bytes()
    config = {
        "paths": {"output_qa_dir": str(tmp_path / "qa")},
        "visual_qa": {
            "allow_blank_pages": True,
            "baseline_dir": str(baseline),
            "minimum_similarity": 1.0,
        },
    }

    draft = _multiformat_service(tmp_path).visual_review(artifact, config, PipelinePolicy())
    strict = _multiformat_service(tmp_path).visual_review(
        artifact, config, PipelinePolicy(PipelineMode.strict)
    )
    release = _multiformat_service(tmp_path).visual_review(
        artifact, config, PipelinePolicy(PipelineMode.release)
    )

    assert draft.ok and "visual.baseline_changed" in draft.detail
    assert not strict.ok and "visual.baseline_changed" in strict.detail
    assert not release.ok and "visual.baseline_changed" in release.detail
    assert baseline_page.read_bytes() == original


def test_docx_visual_stage_does_not_silently_ignore_configured_baseline(tmp_path):
    artifact = _document(tmp_path)
    config = {
        "paths": {"output_qa_dir": str(tmp_path / "qa")},
        "visual_qa": {"baseline_dir": str(tmp_path / "missing-baseline")},
    }
    outcome = _multiformat_service(tmp_path).visual_review(artifact, config, PipelinePolicy(PipelineMode.strict))
    assert not outcome.ok
    assert "baseline" in outcome.detail.lower() or "configured" in outcome.detail.lower()


def test_docx_visual_stage_does_not_silently_ignore_required_previews(tmp_path):
    artifact = _document(tmp_path)
    config = {"visual_qa": {"require_previews": True}}
    outcome = _multiformat_service(tmp_path).visual_review(artifact, config, PipelinePolicy())
    assert not outcome.ok
    assert "preview" in outcome.detail.lower()


def test_pdf_reproducibility_accepts_metadata_only_changes(tmp_path):
    from docs.cli.commands.document_app import _verify_pdf_reproducibility

    artifact = tmp_path / "document.pdf"
    _write_blank_pdf(artifact)
    service = _multiformat_service(tmp_path, pdf_reproducibility=_verify_pdf_reproducibility)

    def rebuild(output):
        output.write_bytes(artifact.read_bytes() + b"\n% different metadata comment\n")
        return output

    outcome = service.reproducibility_check(artifact, PipelinePolicy(), rebuild, tmp_path / "rebuild")
    assert outcome.ok, outcome.detail


def test_docx_visual_stage_uses_rendered_qa_and_keeps_format_audit(tmp_path):
    from docs.application.qa import QaService
    from docs.application.render_verification import RenderVerificationService
    from docs.infrastructure.verification.render_verification_adapter import RenderVerificationAdapter

    class QaPort:
        def render_docx_to_pdf(self, config, docx_path, output_dir):
            pdf = output_dir / "rendered.pdf"
            _write_blank_pdf(pdf)
            return pdf

        def run_documents_audits(self, *args):
            return []

    artifact = _document(tmp_path, heading="Mixed case")
    service = _multiformat_service(tmp_path, qa=QaService(
        QaPort(), FormatAuditService(PythonDocxAuditAdapter()),
        RenderVerificationService(RenderVerificationAdapter()),
    ))
    config = {
        "paths": {"output_qa_dir": str(tmp_path / "qa")},
        "visual_qa": {"baseline_dir": str(tmp_path / "missing"), "allow_blank_pages": True},
    }
    draft = service.visual_review(artifact, config, PipelinePolicy())
    assert draft.ok
    assert "No visual baseline" in draft.detail
    assert "Mixed case" in draft.detail
    assert list((tmp_path / "qa").rglob("*.png"))
    for mode in (PipelineMode.strict, PipelineMode.release):
        outcome = service.visual_review(artifact, config, PipelinePolicy(mode))
        assert not outcome.ok and "No visual baseline" in outcome.detail


def test_pdf_reproducibility_rejects_changed_image_with_identical_geometry(tmp_path):
    import pypdfium2 as pdfium
    from PIL import Image

    from docs.cli.commands.document_app import _verify_pdf_reproducibility

    paths = (tmp_path / "first.pdf", tmp_path / "second.pdf")
    for path, color in zip(paths, ("black", "red"), strict=True):
        with pdfium.PdfDocument.new() as document:
            page = document.new_page(100, 100)
            bitmap = pdfium.PdfBitmap.from_pil(Image.new("RGB", (10, 10), color))
            obj = pdfium.PdfImage.new(document)
            obj.set_bitmap(bitmap)
            obj.set_matrix(pdfium.PdfMatrix(80, 0, 0, 80, 10, 10))
            page.insert_obj(obj)
            page.gen_content()
            document.save(path)
            bitmap.close()
            page.close()
    ok, detail = _verify_pdf_reproducibility(*paths)
    assert not ok, detail
