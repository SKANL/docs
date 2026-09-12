from __future__ import annotations

import hashlib

from docx import Document

from docs.infrastructure.docx.python_docx_assembly_adapter import PythonDocxAssemblyAdapter


def _assemble(tmp_path, variant: str, output_name: str):
    body = Document()
    body.add_paragraph("Body marker")
    body_path = tmp_path / "body.docx"
    body.save(body_path)
    output = tmp_path / output_name
    config = {
        "title": "Native Cover Report",
        "project": {"author": "Ada Lovelace", "institution": "Analytical Academy"},
        "cover": {
            "mode": "generated",
            "variant": variant,
            "slots": {
                "institution": "{{project.institution}}",
                "title": "{{title}}",
                "author": "{{project.author}}",
            },
        },
        "structure": [{"type": "sections"}],
    }
    PythonDocxAssemblyAdapter().assemble(
        config, body_path, output, cover_asset_path=None, embed_front_paths=[], embed_back_paths=[]
    )
    return output


def test_generated_variants_have_distinct_docx_semantics_and_deterministic_bytes(tmp_path):
    outputs = {variant: _assemble(tmp_path, variant, f"{variant}.docx") for variant in (
        "academic", "institutional", "technical", "minimal", "visual", "custom"
    )}

    signatures = {
        variant: (
            len(Document(path).tables),
            tuple(
                (
                    paragraph.text,
                    paragraph.alignment,
                    paragraph.runs[0].font.size.pt if paragraph.runs and paragraph.runs[0].font.size else None,
                    str(paragraph.runs[0].font.color.rgb) if paragraph.runs and paragraph.runs[0].font.color.rgb else None,
                )
                for paragraph in Document(path).paragraphs[:4]
            ),
        )
        for variant, path in outputs.items()
    }

    assert len(set(signatures.values())) == 6
    assert hashlib.sha256(_assemble(tmp_path, "technical", "technical-repeat.docx").read_bytes()).digest() == hashlib.sha256(
        outputs["technical"].read_bytes()
    ).digest()


def test_assemble_without_a_cover_block_preserves_legacy_body_only_output(tmp_path):
    body = Document()
    body.add_paragraph("Body marker")
    body_path = tmp_path / "body.docx"
    body.save(body_path)
    output = tmp_path / "output.docx"

    PythonDocxAssemblyAdapter().assemble(
        {"structure": [{"type": "sections"}]},
        body_path,
        output,
        cover_asset_path=None,
        embed_front_paths=[],
        embed_back_paths=[],
    )

    assert [paragraph.text for paragraph in Document(output).paragraphs if paragraph.text] == ["Body marker"]
