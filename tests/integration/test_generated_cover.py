from __future__ import annotations

import pytest
from docx import Document

from docs.infrastructure.docx.python_docx_assembly_adapter import PythonDocxAssemblyAdapter


def test_assemble_composes_generated_academic_cover_from_declarative_slots(tmp_path):
    body = Document()
    body.add_paragraph("Body marker")
    body_path = tmp_path / "body.docx"
    body.save(body_path)
    output = tmp_path / "output.docx"
    config = {
        "title": "Native Cover Report",
        "project": {"author": "Ada Lovelace", "institution": "Analytical Academy"},
        "cover": {
            "mode": "generated",
            "variant": "academic",
            "slots": {
                "institution": "{{project.institution}}",
                "title": "{{title}}",
                "author": "{{project.author}}",
            },
        },
        "structure": [{"type": "sections"}],
    }

    PythonDocxAssemblyAdapter().assemble(
        config, body_path, output, cover_asset_path=None, embed_front_paths=[], embed_back_paths=[]
    )

    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "Analytical Academy" in text
    assert "Native Cover Report" in text
    assert "Ada Lovelace" in text
    assert text.index("Native Cover Report") < text.index("Body marker")


def test_assemble_keeps_asset_cover_when_legacy_cover_part_is_declared(tmp_path):
    legacy_cover = Document()
    legacy_cover.add_paragraph("LEGACY COVER")
    legacy_path = tmp_path / "legacy.docx"
    legacy_cover.save(legacy_path)
    body = Document()
    body.add_paragraph("Body marker")
    body_path = tmp_path / "body.docx"
    body.save(body_path)
    output = tmp_path / "output.docx"

    PythonDocxAssemblyAdapter().assemble(
        {"structure": [{"type": "cover_from_asset", "asset": "cover"}, {"type": "sections"}]},
        body_path,
        output,
        cover_asset_path=legacy_path,
        embed_front_paths=[],
        embed_back_paths=[],
    )

    assert "LEGACY COVER" in "\n".join(p.text for p in Document(output).paragraphs)


def test_custom_cover_applies_content_page_visual_and_layout_contract(tmp_path):
    body = Document()
    body.add_paragraph("Body marker")
    body_path = tmp_path / "body.docx"
    body.save(body_path)
    output = tmp_path / "output.docx"
    config = {
        "title": "Custom Cover",
        "project": {"owner": "Ada"},
        "format": {
            "cover": {
                "mode": "generated",
                "variant": "custom",
                "content": {"title": "{{title}}", "author": {"path": "project.owner"}},
                "page": {"size": "letter", "margins_cm": {"top": 2}},
                "visual": {"accent_color": "#123456"},
                "layout": {"alignment": "left", "title_size_pt": 28},
            }
        },
        "structure": [{"type": "sections"}],
    }

    PythonDocxAssemblyAdapter().assemble(
        config, body_path, output, cover_asset_path=None, embed_front_paths=[], embed_back_paths=[]
    )

    document = Document(output)
    title = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Custom Cover")
    assert title.alignment == 0  # WD_ALIGN_PARAGRAPH.LEFT
    assert title.runs[0].font.size.pt == 28
    assert "Ada" in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert document.sections[0].top_margin.cm == pytest.approx(2, abs=0.001)


def test_explicit_none_cover_does_not_load_template_or_generate_cover(tmp_path):
    template = Document()
    template.add_paragraph("TEMPLATE COVER MUST NOT APPEAR")
    template_path = tmp_path / "template.docx"
    template.save(template_path)
    body = Document()
    body.add_paragraph("Body marker")
    body_path = tmp_path / "body.docx"
    body.save(body_path)
    output = tmp_path / "output.docx"

    PythonDocxAssemblyAdapter().assemble(
        {
            "paths": {"template_docx": str(template_path)},
            "format": {"cover": {"mode": "none"}},
            "structure": [{"type": "cover_from_template"}, {"type": "sections"}],
        },
        body_path,
        output,
        cover_asset_path=None,
        embed_front_paths=[],
        embed_back_paths=[],
    )

    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "TEMPLATE COVER MUST NOT APPEAR" not in text
    assert "Body marker" in text
