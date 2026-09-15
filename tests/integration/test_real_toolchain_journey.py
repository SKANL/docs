from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document

from docs.domain.ports.visual_renderer_port import VisualSpec
from docs.infrastructure.docx.libreoffice_qa_adapter import resolve_libreoffice_executable
from docs.infrastructure.docx.tool_resolver_adapter import SystemToolResolverAdapter
from docs.infrastructure.ingest.pandoc_ingest_adapter import PandocIngestAdapter
from docs.infrastructure.visuals.mermaid_svg_renderer import MermaidSvgRenderer
from docs.infrastructure.visuals.resvg_rasterizer_adapter import ResvgRasterizerAdapter

_RESOLVER = SystemToolResolverAdapter()
_REQUIRED_TOOLS = {
    "pandoc": _RESOLVER.resolve_pandoc({}),
    "libreoffice": resolve_libreoffice_executable({}),
    "mmdc": _RESOLVER.resolve_mmdc({}),
    "resvg": _RESOLVER.resolve_resvg({}),
    "pdfinfo": shutil.which("pdfinfo"),
    "pdftoppm": shutil.which("pdftoppm"),
}
_MISSING_TOOLS = tuple(name for name, path in _REQUIRED_TOOLS.items() if not path)


def test_traceability_distinguishes_real_toolchain_proof_from_unavailable_capabilities() -> None:
    traceability = json.loads(
        (Path(__file__).parents[2] / "docs" / "traceability.json").read_text(
            encoding="utf-8"
        )
    )

    journey = traceability["real_toolchain_journey"]
    assert journey["status"] == "proved-when-capable"
    assert journey["test"] == "tests/integration/test_v2_real_toolchain_journey.py::test_v2_real_toolchain_journey"
    assert journey["proved"] == ["pandoc", "libreoffice", "poppler", "mermaid", "resvg"]
    assert journey["unavailable"] == "controlled skip with capability names in pytest output"


@pytest.mark.skipif(bool(_MISSING_TOOLS), reason=f"external toolchain unavailable: {', '.join(_MISSING_TOOLS)}")
def test_v2_real_toolchain_journey(tmp_path: Path) -> None:
    """Exercise the real external adapters in one bounded v2 toolchain journey."""
    source_docx = tmp_path / "source.docx"
    document = Document()
    document.add_heading("External journey", level=1)
    document.add_paragraph("Pandoc, Mermaid, resvg, LibreOffice, and Poppler proof.")
    document.save(str(source_docx))

    ingested_dir = tmp_path / "ingested"
    ingested_dir.mkdir()
    ingested = PandocIngestAdapter(_RESOLVER).ingest(source_docx, ingested_dir, "docx")
    assert "External journey" in ingested.read_text(encoding="utf-8")

    svg = MermaidSvgRenderer(_RESOLVER, scratch_root=tmp_path).render(
        VisualSpec(label="journey", type="mermaid", source="graph TD; A-->B;")
    )
    svg_path = tmp_path / "journey.svg"
    svg_path.write_text(svg, encoding="utf-8")
    png_path = tmp_path / "journey.png"
    ResvgRasterizerAdapter(_RESOLVER).rasterize(svg_path, png_path)
    assert png_path.stat().st_size > 0

    pdf_path = tmp_path / "pdf"
    pdf_path.mkdir()
    from docs.infrastructure.docx.libreoffice_qa_adapter import LibreOfficeQaAdapter

    rendered_pdf = LibreOfficeQaAdapter().render_docx_to_pdf({}, source_docx, pdf_path)
    assert rendered_pdf.stat().st_size > 0

    pdfinfo = _REQUIRED_TOOLS["pdfinfo"]
    pdftoppm = _REQUIRED_TOOLS["pdftoppm"]
    assert pdfinfo is not None
    assert pdftoppm is not None
    info = subprocess.run([pdfinfo, str(rendered_pdf)], check=True, capture_output=True, text=True)
    preview = tmp_path / "preview"
    preview.mkdir()
    subprocess.run([pdftoppm, "-png", str(rendered_pdf), str(preview / "page")], check=True)
    assert b"Pages:" in info.stdout.encode()
    assert list(preview.glob("page-*.png"))

