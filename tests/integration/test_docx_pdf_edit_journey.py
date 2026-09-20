from __future__ import annotations

import base64
import hashlib
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches

from docs.application.docx_import import DocxImportService
from docs.application.evidence_passport import EvidencePassportService
from docs.application.pdf_render import PdfRendererAdapter
from docs.application.provenance import ProvenanceLedger
from docs.domain.ports.pdf_text_edit_port import BlockReplacement
from docs.domain.text_fitting import FittedText
from docs.infrastructure.docx.libreoffice_qa_adapter import (
    LibreOfficeQaAdapter,
    resolve_libreoffice_executable,
)
from docs.infrastructure.docx.python_docx_assembly_adapter import PythonDocxAssemblyAdapter
from docs.infrastructure.docx.python_docx_audit_adapter import PythonDocxAuditAdapter
from docs.infrastructure.docx.tool_resolver_adapter import SystemToolResolverAdapter
from docs.infrastructure.pdf.pypdfium2_text_edit_adapter import Pypdfium2TextEditAdapter
from docs.infrastructure.persistence.evidence_passport_store import FileEvidencePassportStore

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUB"
    "AScY42YAAAAASUVORK5CYII="
)


def _make_source_docx(path: Path) -> None:
    image = path.with_suffix(".png")
    image.write_bytes(_PNG)
    document = Document()
    document.add_heading("IMPORTED JOURNEY", level=1)
    document.add_paragraph("UNTOUCHED SOURCE SENTENCE.")
    document.add_picture(str(image), width=Inches(0.25))
    document.add_paragraph("Figure 1. Preserved source asset.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "UNTOUCHED TABLE"
    table.cell(0, 1).text = "VALUE"
    document.add_paragraph("EDITABLE SENTENCE THAT WILL BE REPLACED.")
    document.save(path)


def _pdf_geometry(path: Path) -> tuple[tuple[float, float], ...]:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(path))
    try:
        return tuple((page.get_width(), page.get_height()) for page in document)
    finally:
        document.close()


def test_docx_import_to_pdf_edit_journey_is_deterministic_and_bounded(tmp_path: Path) -> None:
    pandoc = SystemToolResolverAdapter().resolve_pandoc({})
    libreoffice = resolve_libreoffice_executable({})
    if not pandoc or not libreoffice:
        pytest.skip("real DOCX/PDF journey requires pandoc and LibreOffice")

    source = tmp_path / "source.docx"
    _make_source_docx(source)
    imported = tmp_path / "imported"
    report = DocxImportService().import_file(source, imported)
    repeat = tmp_path / "repeat"
    repeat_report = DocxImportService().import_file(source, repeat)

    assert report == repeat_report
    assert report.source_hash == hashlib.sha256(source.read_bytes()).hexdigest()
    assert (imported / "source.docx").read_bytes() == source.read_bytes()
    assert (imported / "assets" / "image-1.png").read_bytes() == _PNG
    assert "source:exact-bytes" in report.preserved
    assert "asset:embedded-image-bytes" in report.preserved
    assert "empty-paragraph" in report.dropped
    assert report.unsupported == ()
    assert sorted(path.relative_to(imported).as_posix() for path in imported.rglob("*") if path.is_file()) == sorted(
        path.relative_to(repeat).as_posix() for path in repeat.rglob("*") if path.is_file()
    )

    section = imported / "editable-section.md"
    section.write_text(
        (imported / "document.md")
        .read_text(encoding="utf-8")
        .replace("EDITABLE SENTENCE THAT WILL BE REPLACED.", "EDITED SECTION SENTENCE."),
        encoding="utf-8",
        newline="\n",
    )
    assert "UNTOUCHED SOURCE SENTENCE." in section.read_text(encoding="utf-8")
    assert (imported / "assets" / "image-1.png").read_bytes() == _PNG

    rendered = tmp_path / "rendered.docx"
    PythonDocxAssemblyAdapter().render_pandoc(str(pandoc), [section], rendered)
    rendered_document = Document(rendered)
    rendered_text = "\n".join(paragraph.text for paragraph in rendered_document.paragraphs)
    assert "EDITED SECTION SENTENCE." in rendered_text
    assert "UNTOUCHED SOURCE SENTENCE." in rendered_text
    assert any("UNTOUCHED TABLE" in cell.text for table in rendered_document.tables for row in table.rows for cell in row.cells)
    assert rendered_document.inline_shapes
    assert not [issue for issue in PythonDocxAuditAdapter().audit(rendered, {}, strict=False) if issue.severity == "error"]

    pdf_dir = tmp_path / "pdf"
    pdf_dir.mkdir()
    rendered_pdf = LibreOfficeQaAdapter().render_docx_to_pdf({}, rendered, pdf_dir)
    before_geometry = _pdf_geometry(rendered_pdf)
    adapter = Pypdfium2TextEditAdapter()
    target = next(run for run in adapter.read_runs(rendered_pdf) if run.text == "EDITED SECTION SENTENCE.")
    edited_pdf = tmp_path / "edited.pdf"
    replacement = BlockReplacement(
        page=target.page,
        remove=[target],
        fitted=FittedText(["PDF EDIT."], target.font_size, False),
        x=target.x,
        top=target.top,
        baseline=target.y,
        right=target.right,
    )
    write_report = adapter.write_blocks(
        rendered_pdf,
        edited_pdf,
        [replacement],
    )
    assert write_report.verification_diagnostics == ()
    assert write_report.blocks_unsafe == 0
    after_geometry = _pdf_geometry(edited_pdf)
    assert len(after_geometry) == len(before_geometry)
    for after_page, before_page in zip(after_geometry, before_geometry, strict=True):
        assert after_page == pytest.approx(before_page)
    edited_text = {run.text for run in adapter.read_runs(edited_pdf)}
    assert "PDF EDIT." in edited_text
    assert "UNTOUCHED SOURCE SENTENCE." in edited_text
    assert "UNTOUCHED TABLE" in " ".join(edited_text)
    repeated_pdf = tmp_path / "edited-repeat.pdf"
    repeated_report = adapter.write_blocks(rendered_pdf, repeated_pdf, [replacement])
    assert repeated_report == write_report
    assert repeated_pdf.read_bytes() == edited_pdf.read_bytes()

    run_id = "docx-pdf-edit-journey"
    ledger = ProvenanceLedger(tmp_path / "provenance.json", trusted_root=tmp_path)
    run = ledger.record_run(run_id, inputs=(source, section), outputs=(rendered, edited_pdf))
    attestation = {"run_id": run_id, "outputs": run["outputs"], "pdf_write": write_report.blocks_written}
    ledger.record_attestation(run_id, attestation)
    assert ledger.verify_attestation(run_id, attestation)

    passport_service = EvidencePassportService(FileEvidencePassportStore(tmp_path / "passports"))
    passport = passport_service.finalize(
        run_id,
        (
            {"source_hash": report.source_hash, "import_report": report.__dict__},
            {"artifact": str(edited_pdf), "geometry": _pdf_geometry(edited_pdf), "write_report": attestation},
        ),
    )
    assert passport == passport_service.finalize(
        run_id,
        (
            {"source_hash": report.source_hash, "import_report": report.__dict__},
            {"artifact": str(edited_pdf), "geometry": _pdf_geometry(edited_pdf), "write_report": attestation},
        ),
    )


def test_docx_pdf_journey_reports_missing_libreoffice_without_partial_pdf(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    source = tmp_path / "source.docx"
    _make_source_docx(source)

    class _ExistingDocxRenderer:
        output_format = "docx"

        def build(self, _doc_id: str, _config: dict[str, object]) -> Path:
            return source

    monkeypatch.setattr(
        "docs.infrastructure.docx.libreoffice_qa_adapter.resolve_libreoffice_executable",
        lambda _paths: None,
    )
    output_dir = tmp_path / "output"
    output_dir.mkdir()
    result = PdfRendererAdapter(
        _ExistingDocxRenderer(),
        LibreOfficeQaAdapter(),
    ).build("journey", {"paths": {"output_draft_dir": str(output_dir)}})

    assert result is None
    assert not list(output_dir.glob("*.pdf"))
    captured = capsys.readouterr()
    assert "WARN" in captured.err
    assert "PDF" in captured.err
