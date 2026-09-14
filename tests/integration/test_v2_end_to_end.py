from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree as ET

import pytest
from PIL import Image
from typer.testing import CliRunner

from docs.application.asset import AssetService
from docs.application.documents import DocumentService
from docs.application.docx_assembly import DocxRendererAdapter
from docs.application.generate_visuals import GenerateVisualsService
from docs.application.pipeline_service_v2 import FULL_STAGE_IDS
from docs.cli.main import app
from docs.domain.models.template import Template
from docs.domain.workspace import Workspace
from docs.infrastructure.docx.python_docx_assembly_adapter import PythonDocxAssemblyAdapter
from docs.infrastructure.docx.python_docx_image_metadata_adapter import PythonDocxImageMetadataAdapter
from docs.infrastructure.docx.tool_resolver_adapter import SystemToolResolverAdapter
from docs.infrastructure.ingest.atomic_file_adapter import AtomicFileAdapter
from docs.infrastructure.ingest.filesystem_ingest_artifact_writer import FilesystemIngestArtifactWriter
from docs.infrastructure.ingest.md_normalize_adapter import MdNormalizeAdapter
from docs.infrastructure.persistence.filesystem_asset_repository import FilesystemAssetRepository
from docs.infrastructure.persistence.json_repository import JsonDocumentRepository
from docs.infrastructure.visuals.chart_svg_renderer import ChartSvgRenderer


class _JourneyIngest:
    def ingest_inbox(self, inbox_dir, sections_dir, assets_dir=None):
        del assets_dir
        source = next(Path(inbox_dir).glob("*.md"))
        output = Path(sections_dir) / "ingested" / "brief-md-journey.md"
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(
            "---\n{\"source\": \"journey\"}\n---\n" + source.read_text(encoding="utf-8"),
            encoding="utf-8",
        )
        return {"processed": 1, "files": [{"status": "converted", "path": str(output)}]}


class _JourneyRenderer:
    output_format = "docx"
    required_capabilities: tuple[str, ...] = ()
    version = "journey-renderer-1"

    def __init__(self) -> None:
        self.calls: list[Path] = []

    def build(self, doc_id: str, config: dict[str, object], output: Path | None = None) -> Path:
        del config
        assert output is not None
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(f"DOCX:{doc_id}:journey".encode())
        self.calls.append(output)
        return output


class _JourneyAudit:
    def audit_format(self, path: Path, config: dict[str, object], strict: bool = False):
        del path, config, strict
        from docs.domain.review import ReviewResult

        return ReviewResult()


class _JourneyQa:
    def qa_docx(self, config: dict[str, object], path: Path, strict: bool = False) -> Path:
        del path, strict
        target = Path(config["paths"]["output_qa_dir"])
        target.mkdir(parents=True, exist_ok=True)
        (target / "qa-report.md").write_text("verified", encoding="utf-8")
        return target


class _JourneyReview:
    def review_document(self, *args, **kwargs):
        from docs.domain.review import ReviewResult

        return ReviewResult()


class _JourneySvgRasterizer:
    """Existing internal rasterizer seam, pinned to a deterministic PNG for this journey."""

    def rasterize(self, svg_path: Path, png_path: Path) -> None:
        del svg_path
        png_path.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (64, 32), color=(10, 20, 30)).save(png_path, format="PNG")


def _journey_deps(tmp_path: Path):
    documents_dir = tmp_path / "documents"
    templates_dir = tmp_path / "templates"
    workspace = Workspace(documents_dir, templates_dir)
    repository = JsonDocumentRepository(workspace)
    documents = DocumentService(repository, workspace, clock=lambda: "2026-01-01T00:00:00")
    template = Template(type="journey", title="Journey", structure=[{"type": "sections"}])
    templates_dir.mkdir(parents=True)
    (templates_dir / "journey.json").write_text(template.model_dump_json(), encoding="utf-8")
    renderer = _JourneyRenderer()
    audit = _JourneyAudit()
    qa = _JourneyQa()
    review = _JourneyReview()

    def resolve_context(doc_id: str = ""):
        selected = doc_id or repository.active_id()
        root = workspace.doc_root(selected)
        config = {
            "output": {"format": "docx"},
            "structure": [{"type": "sections"}],
            "paths": {
                "inbox_dir": str(root / "inbox"),
                "sections_dir": str(root / "sections"),
                "assets_dir": str(root / "assets"),
                "runs_dir": str(root / "runs"),
                "output_qa_dir": str(root / "output" / "qa"),
                "output_draft_dir": str(root / "output" / "draft"),
            },
        }
        return SimpleNamespace(
            doc_id=selected,
            config=config,
            template=repository.load_template("journey"),
        )

    def stage() -> tuple[bool, str]:
        return True, "journey"
    pipeline = SimpleNamespace(
        generate_visuals=stage,
        compose_cover=stage,
        build_html=stage,
        build_pdf=stage,
        structural_audit=stage,
        accessibility_review=stage,
        visual_review=stage,
        reproducibility_check=stage,
        evidence_review=stage,
        consistency_review=stage,
        rules_manifest_state=lambda config: (True, 1),
    )
    return SimpleNamespace(
        workspace=workspace,
        document_repository=repository,
        documents=documents,
        ingest=_JourneyIngest(),
        markdown_normalizer=MdNormalizeAdapter(),
        atomic_file_writer=AtomicFileAdapter(),
        resolve_context=resolve_context,
        resolve_renderer=lambda config: renderer,
        renderers={"docx": renderer},
        renderer=renderer,
        format_audit=audit,
        qa=qa,
        review=review,
        pipeline=pipeline,
        v2_compatibility=pipeline,
    )


def invoke(runner: CliRunner, *args: str):
    result = runner.invoke(app, ["document", *args, "--json"])
    assert result.exit_code == 0, result.stdout
    return json.loads(result.stdout)


def test_v2_public_journey_creates_prepares_builds_verifies_packages_and_publishes(
    monkeypatch, tmp_path: Path
) -> None:
    deps = _journey_deps(tmp_path)
    monkeypatch.setattr("docs.cli.main.Deps", lambda: deps)
    runner = CliRunner()

    created = invoke(runner, "create", "journey", "--template", "journey", "--title", "Journey")
    root = deps.workspace.doc_root("journey")
    (root / "inbox" / "source.md").write_text("A source claim.\n", encoding="utf-8")
    assert created["document_id"] == "journey"
    assert (root / "document.json").is_file()

    ingested = invoke(runner, "ingest")
    prepared = invoke(runner, "prepare")
    built = invoke(runner, "build")
    verified = invoke(runner, "verify")

    assert ingested["stages"][0]["name"] == "ingest-sources"
    assert [stage["name"] for stage in prepared["stages"]] == [
        "ingest-sources",
        "normalize-sources",
        "compile-structure",
    ]
    assert built["report"]["succeeded"] is True
    assert verified["report"]["succeeded"] is True

    artifact = root / "output" / "v2" / "journey.docx"
    package = root / "output" / "release" / "journey-external.zip"
    packaged = invoke(runner, "package", str(artifact.parent), str(package))
    published = invoke(
        runner,
        "publish",
        str(artifact),
        str(root / "published" / "journey.docx"),
    )

    assert packaged["path"] == str(package.resolve())
    assert published["published"] is True
    assert (root / "published" / "journey.docx").read_bytes() == artifact.read_bytes()
    assert all(deps.renderer.calls)
    assert (root / "runs" / "v2-prepare.json").is_file()
    assert set(FULL_STAGE_IDS) >= {
        "ingest-sources",
        "normalize-sources",
        "compile-structure",
        "build-docx",
        "editorial-review",
        "package-release",
        "publish-draft",
    }


def test_v2_release_command_composes_creation_and_verified_build(monkeypatch, tmp_path: Path) -> None:
    deps = _journey_deps(tmp_path)
    monkeypatch.setattr("docs.cli.main.Deps", lambda: deps)
    runner = CliRunner()
    created = runner.invoke(
        app,
        ["document", "create", "one-shot", "--template", "journey", "--title", "One-shot", "--json"],
    )
    assert created.exit_code == 0, created.stdout
    (deps.workspace.doc_root("one-shot") / "inbox" / "source.md").write_text(
        "A source claim.\n", encoding="utf-8"
    )

    result = runner.invoke(
        app,
        [
            "document",
            "release",
            "--json",
        ],
    )

    assert result.exit_code == 0, result.stdout
    payload = json.loads(result.stdout)
    assert payload["report"]["succeeded"] is True
    root = deps.workspace.doc_root("one-shot")
    assert (root / "document.json").is_file()
    assert (root / "output" / "v2" / "one-shot.docx").is_file()
    assert (root / "output" / "release" / "one-shot.zip").is_file()


def test_v2_builtin_template_journey_produces_artifacts_and_provenance(
    monkeypatch, tmp_path: Path
) -> None:
    """The public journey must work in an isolated workspace with a shipped template."""
    deps = _journey_deps(tmp_path)
    fixture = Path(__file__).parents[1] / "fixtures" / "templates" / "documento-generico.json"
    (deps.workspace.templates_dir / "documento-generico.json").write_bytes(fixture.read_bytes())
    monkeypatch.setattr("docs.cli.main.Deps", lambda: deps)
    runner = CliRunner()

    created = invoke(
        runner,
        "create",
        "canonical",
        "--template",
        "documento-generico",
        "--title",
        "Canonical journey",
    )
    root = deps.workspace.doc_root("canonical")
    (root / "inbox" / "source.md").write_text("A source claim.\n", encoding="utf-8")

    ingested = invoke(runner, "ingest")
    prepared = invoke(runner, "prepare")
    built = invoke(runner, "build")
    verified = invoke(runner, "verify")

    artifact = root / "output" / "v2" / "canonical.docx"
    inspected = invoke(runner, "inspect", str(artifact))
    package = root / "output" / "release" / "canonical.zip"
    packaged = invoke(runner, "package", str(artifact.parent), str(package))
    destination = root / "published" / "canonical.docx"
    published = invoke(runner, "publish", str(artifact), str(destination))

    assert created["document_id"] == "canonical"
    assert ingested["stages"][0]["name"] == "ingest-sources"
    assert [stage["name"] for stage in prepared["stages"]] == [
        "ingest-sources",
        "normalize-sources",
        "compile-structure",
    ]
    assert built["report"]["succeeded"] is True
    assert verified["report"]["succeeded"] is True
    assert artifact.read_bytes() == b"DOCX:canonical:journey"
    assert inspected["sha256"] == hashlib.sha256(artifact.read_bytes()).hexdigest()
    assert packaged["path"] == str(package.resolve())
    with zipfile.ZipFile(package) as archive:
        assert "canonical.docx" in archive.namelist()

    manifest_path = artifact.with_suffix(".docx.manifest.json")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["document_id"] == "canonical"
    assert manifest["verification"]["passed"] is True
    assert manifest["artifacts"][0]["sha256"] == inspected["sha256"]
    assert manifest["provenance_run"].startswith("cli-build-docx-")
    assert manifest["provenance_run"] in (root / "runs" / "v2-provenance.json").read_text(encoding="utf-8")
    assert published["published"] is True
    assert destination.read_bytes() == artifact.read_bytes()
    assert destination.with_suffix(".docx.manifest.json").read_bytes() == manifest_path.read_bytes()


def test_v2_stage_traceability_declares_every_runtime_stage() -> None:
    path = Path(__file__).parents[2] / "docs" / "migration-v2-traceability.json"
    traceability = json.loads(path.read_text(encoding="utf-8"))
    declared = {entry["stage"] for entry in traceability["pipeline_stages"]}
    assert declared == set(FULL_STAGE_IDS)
    assert traceability["journey"]["status"] in {"covered", "closest-real-journey"}


def test_v2_journey_executes_generated_cover_and_visual_spec_fixture(monkeypatch, tmp_path: Path) -> None:
    fixture = json.loads(
        (Path(__file__).parents[1] / "fixtures" / "v2" / "generated-cover-visual-journey.json").read_text(
            encoding="utf-8"
        )
    )
    deps = _journey_deps(tmp_path)
    tool_resolver = SystemToolResolverAdapter()
    real_renderer = DocxRendererAdapter(
        PythonDocxAssemblyAdapter(),
        AssetService(FilesystemAssetRepository(), deps.workspace),
        tool_resolver,
    )
    deps.renderers = {"docx": real_renderer}
    deps.resolve_renderer = lambda config: real_renderer
    deps.generate_visuals_service = GenerateVisualsService(
        {"chart": ChartSvgRenderer()},
        _JourneySvgRasterizer(),
        image_metadata=PythonDocxImageMetadataAdapter(),
        writer=FilesystemIngestArtifactWriter(),
    )
    # The v2 application provider owns these stages; no CLI compatibility
    # hook is required for the real visual/cover journey.
    deps.pipeline.generate_visuals = None
    deps.pipeline.compose_cover = None
    original_resolve_context = deps.resolve_context

    def resolve_context(doc_id: str = ""):
        resolved = original_resolve_context(doc_id)
        resolved.config.update(fixture["document"])
        resolved.config["title"] = fixture["title"]
        return resolved

    deps.resolve_context = resolve_context
    monkeypatch.setattr("docs.cli.main.Deps", lambda: deps)
    runner = CliRunner()

    created = invoke(runner, "create", "generated", "--template", "journey", "--title", fixture["title"])
    root = deps.workspace.doc_root("generated")
    document = json.loads((root / "document.json").read_text(encoding="utf-8"))
    document.update(fixture["document"])
    (root / "document.json").write_text(json.dumps(document), encoding="utf-8")
    (root / "inbox" / "source.md").write_text(fixture["source_markdown"], encoding="utf-8")
    (root / "sections" / "visual-specs.json").write_text(
        json.dumps(fixture["visual_specs"]), encoding="utf-8"
    )

    invoke(runner, "ingest")
    invoke(runner, "prepare")
    ingested_section = root / "sections" / "ingested" / "brief-md-journey.md"
    (root / "sections" / "001-brief-md-journey.md").write_bytes(ingested_section.read_bytes())
    authored_markdown = (root / "sections" / "001-brief-md-journey.md").read_text(encoding="utf-8")
    assert fixture["source_markdown"].splitlines() == [
        "# Journey body",
        "",
        "Journey body with a generated visual.",
        "",
        "[[figure:journey-chart]] Revenue evidence.",
    ]
    authored_lines = authored_markdown.splitlines()
    assert "# Journey body" in authored_lines
    assert "Journey body with a generated visual." in authored_lines
    assert "[[figure:journey-chart]] Revenue evidence." in authored_lines
    assert r"\n" not in authored_markdown
    first = invoke(runner, "build")
    artifact = root / "output" / "v2" / "generated.docx"

    assert created["document_id"] == "generated"
    assert artifact.is_file()
    figure_files = sorted((root / "assets" / "figures").glob("*.png"))
    assert figure_files
    catalog = json.loads((root / "sections" / "figure-catalog.json").read_text(encoding="utf-8"))
    bindings = json.loads((root / "sections" / "figure-bindings.json").read_text(encoding="utf-8"))
    assert catalog["figures"]
    bound_id = bindings["bindings"][fixture["visual_specs"][0]["label"]]
    assert bound_id.startswith("fig-")
    bound_catalog_row = next(row for row in catalog["figures"] if row["id"] == bound_id)
    bound_media_name = Path(bound_catalog_row["origin_relative_path"]).name

    stage_results = first["report"]["execution"]["results"]
    cover_stage = next(result for result in stage_results if result["stage"] == "compose-cover")
    visuals_stage = next(result for result in stage_results if result["stage"] == "generate-visuals")
    assert cover_stage["ok"] is True
    assert cover_stage["outcome"] == "succeeded"
    assert cover_stage["errors"] == []
    assert visuals_stage["ok"] is True
    assert visuals_stage["outcome"] == "succeeded"
    assert visuals_stage["errors"] == []

    from docx import Document

    paragraphs = [paragraph.text for paragraph in Document(artifact).paragraphs]
    document_text = "\n".join(paragraphs)
    assert "Generated Journey Cover" in document_text
    assert "Journey body" in document_text
    assert "generated visual." in document_text
    assert "Figura 1. Revenue evidence." in document_text

    # The v2 registry deliberately keeps generation after the first build;
    # the next build consumes the generated catalog/binding artifacts.
    second = invoke(runner, "build")
    assert second["report"]["succeeded"] is True
    second_bytes = artifact.read_bytes()
    second_sha = hashlib.sha256(second_bytes).hexdigest()
    reopened = Document(artifact)
    assert any("Generated Journey Cover" in paragraph.text for paragraph in reopened.paragraphs)
    with zipfile.ZipFile(artifact) as archive:
        bound_media_bytes = (root / "assets" / "figures" / bound_media_name).read_bytes()
        rels = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
        image_relationships = {
            relationship.attrib["Id"]: relationship.attrib["Target"]
            for relationship in rels
            if relationship.attrib.get("Type", "").endswith("/image")
        }
        media_targets = {
            f"word/{target}": relationship_id
            for relationship_id, target in image_relationships.items()
            if target.startswith("media/")
        }
        assert media_targets
        matching_media = {
            media_name: relationship_id
            for media_name, relationship_id in media_targets.items()
            if archive.read(media_name) == bound_media_bytes
        }
        assert matching_media
        document_xml = ET.fromstring(archive.read("word/document.xml"))
        embed_ids = {
            blip.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"]
            for blip in document_xml.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
        }
        assert set(matching_media.values()).issubset(embed_ids)

    third = invoke(runner, "build")
    assert third["report"]["succeeded"] is True
    assert artifact.read_bytes() == second_bytes
    assert hashlib.sha256(artifact.read_bytes()).hexdigest() == second_sha


_BUILTIN_TEMPLATE_CASES = {
    "reporte-estadia-tic": {
        "structure_types": ["cover_from_asset", "blank_page", "fixed_text_page", "sections"],
        "required_contract": "resumen",
        "context_topic": "alumno",
    },
    "technical-report-srs": {
        "structure_types": ["fixed_text_page", "toc", "sections"],
        "required_contract": "requirements",
        "context_topic": "project",
    },
    "documento-generico": {
        "structure_types": ["fixed_text_page", "sections"],
        "required_contract": "introduccion",
        "context_topic": "documento",
    },
}


@pytest.mark.parametrize("template_name", _BUILTIN_TEMPLATE_CASES)
def test_v2_build_and_verify_honor_each_builtin_template_contract(
    monkeypatch, tmp_path: Path, template_name: str
) -> None:
    deps = _journey_deps(tmp_path)
    fixture_dir = Path(__file__).parents[1] / "fixtures" / "templates"
    fixture = fixture_dir / f"{template_name}.json"
    (deps.workspace.templates_dir / fixture.name).write_bytes(fixture.read_bytes())

    repository = deps.document_repository
    original_resolve_context = deps.resolve_context

    def resolve_context(doc_id: str = ""):
        selected = doc_id or repository.active_id()
        document = repository.read_document(selected)
        template = repository.load_template(document.template)
        resolved = original_resolve_context(selected)
        config = template.model_dump()
        config["paths"] = resolved.config["paths"]
        config["doc_id"] = selected
        return SimpleNamespace(doc_id=selected, config=config, template=template)

    deps.resolve_context = resolve_context
    monkeypatch.setattr("docs.cli.main.Deps", lambda: deps)
    runner = CliRunner()

    created = invoke(
        runner,
        "create",
        "matrix",
        "--template",
        template_name,
        "--title",
        "Matrix",
    )
    root = deps.workspace.doc_root("matrix")
    (root / "inbox" / "source.md").write_text("A source claim.\n", encoding="utf-8")
    invoke(runner, "ingest")
    invoke(runner, "prepare")
    built = invoke(runner, "build")
    verified = invoke(runner, "verify")

    template = repository.load_template(template_name)
    expected = _BUILTIN_TEMPLATE_CASES[template_name]
    assert created["template"] == template_name
    assert [part["type"] for part in template.structure] == expected["structure_types"]
    assert expected["required_contract"] in template.section_contracts
    assert any(topic.id == expected["context_topic"] for topic in template.context_schema.topics)
    assert built["report"]["succeeded"] is True
    assert verified["report"]["succeeded"] is True

    artifact = root / "output" / "v2" / "matrix.docx"
    manifest_path = artifact.with_suffix(".docx.manifest.json")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert artifact.is_file()
    assert manifest["schema"] == "docs.build/v2"
    assert manifest["document_id"] == "matrix"
    assert manifest["template_hash"]
    assert manifest["verification"]["passed"] is True
    assert manifest["artifacts"][0]["path"] == str(artifact.resolve())
    assert manifest["provenance_run"].startswith("cli-build-docx-")
    assert (root / "runs" / "v2-provenance.json").is_file()
