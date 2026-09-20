from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
from dataclasses import dataclass
from importlib import import_module
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile


@dataclass(frozen=True)
class ImportReport:
    preserved: tuple[str, ...]
    normalized: tuple[str, ...]
    dropped: tuple[str, ...]
    unsupported: tuple[str, ...]
    source_hash: str


class DocxImportService:
    """Import the readable DOCX structure without pretending to round-trip it."""

    def import_file(self, path: str | Path, output_dir: str | Path) -> ImportReport:
        source = Path(path)
        output = Path(output_dir)
        source_bytes = source.read_bytes()
        source_hash = hashlib.sha256(source_bytes).hexdigest()
        try:
            docx = import_module("docx")
        except ModuleNotFoundError as exc:
            raise RuntimeError("python-docx is required for DOCX import") from exc

        try:
            with ZipFile(source) as archive:
                names = set(archive.namelist())
        except (BadZipFile, OSError) as exc:
            raise ValueError(f"Malformed DOCX: {source}") from exc

        try:
            document = docx.Document(str(source))
        except Exception as exc:
            raise ValueError(f"Malformed DOCX: {source}") from exc

        preserved = ["source:exact-bytes", "body:block-order"]
        normalized: list[str] = []
        dropped: list[str] = []
        unsupported = [f"ooxml-part:{name}" for name in sorted(names) if self._unsupported_part(name)]
        markdown, assets = self._markdown(document, normalized, dropped)
        if assets:
            preserved.append("asset:embedded-image-bytes")
        report = ImportReport(tuple(preserved), tuple(normalized), tuple(dropped), tuple(unsupported), source_hash)
        self._publish(output, source_bytes, markdown, assets, report)
        return report

    @staticmethod
    def _unsupported_part(name: str) -> bool:
        return name.startswith("word/") and name.endswith(".xml") and name.split("/", 1)[1].split(".", 1)[0] in {
            "footnotes", "endnotes", "comments", "customXml", "glossary",
        }

    def _markdown(self, document: Any, normalized: list[str], dropped: list[str]) -> tuple[str, dict[str, bytes]]:
        from docx.table import Table

        lines = ["---"]
        properties = document.core_properties
        for key in ("title", "author", "subject", "keywords"):
            value = getattr(properties, key, "")
            if value:
                lines.append(f'{key}: {json.dumps(value, ensure_ascii=False)}')
        lines += ["---", ""]
        assets: dict[str, bytes] = {}
        image_number = 0
        seen_image_rel_ids: set[str] = set()

        def append_images(paragraph: Any) -> None:
            nonlocal image_number
            for shape in paragraph._p.xpath('.//*[local-name()="blip"]'):
                rel_id = shape.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                if not rel_id or rel_id in seen_image_rel_ids or rel_id not in document.part.rels:
                    continue
                seen_image_rel_ids.add(rel_id)
                image_number += 1
                part = document.part.rels[rel_id].target_part
                filename = f"image-{image_number}{Path(part.partname).suffix.lower()}"
                assets[filename] = part.blob
                lines.extend([f"![Image {image_number}](assets/{filename})", ""])

        for block in self._iter_body_blocks(document):
            if isinstance(block, Table):
                rows = [
                    [cell.text.replace("|", "\\|").replace("\n", " ") for cell in row.cells]
                    for row in block.rows
                ]
                if rows:
                    lines.extend(
                        [
                            "| " + " | ".join(rows[0]) + " |",
                            "| " + " | ".join("---" for _ in rows[0]) + " |",
                        ]
                    )
                    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
                    lines.append("")
                for row in block.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            append_images(paragraph)
                continue

            paragraph = block
            text = paragraph.text
            style_name = getattr(paragraph.style, "name", "")
            if not text:
                dropped.append("empty-paragraph")
                continue
            if style_name and style_name.startswith("Heading "):
                try:
                    level = int(style_name.rsplit(" ", 1)[1])
                except ValueError:
                    level = 1
                lines.extend([f"{'#' * level} {text}", ""])
            elif style_name and style_name not in {"Normal", "Caption"}:
                normalized.append(f"paragraph-style:{style_name}")
                lines.extend([f"<!-- style: {style_name} -->", text, ""])
            elif style_name == "Caption":
                lines.extend([f"*{text}*", ""])
            else:
                lines.extend([text, ""])
            append_images(paragraph)

        for rel_id, relationship in sorted(document.part.rels.items()):
            if "image" not in relationship.reltype:
                continue
            if rel_id in seen_image_rel_ids:
                continue
            seen_image_rel_ids.add(rel_id)
            image_number += 1
            part = relationship.target_part
            filename = f"image-{image_number}{Path(part.partname).suffix.lower()}"
            if filename not in assets:
                assets[filename] = part.blob
                lines.extend([f"![Image {image_number}](assets/{filename})", ""])

        for title, sections in (("Headers", document.sections), ("Footers", document.sections)):
            values = [getattr(section, "header" if title == "Headers" else "footer").paragraphs for section in sections]
            content = [(index, p.text) for index, paragraphs in enumerate(values, 1) for p in paragraphs if p.text]
            if content:
                lines.extend([f"## {title}", ""])
                for index, text in content:
                    lines.extend([f"### Section {index}", "", text, ""])

        return "\n".join(lines).rstrip() + "\n", assets

    @staticmethod
    def _iter_body_blocks(document: Any):
        """Yield paragraphs and tables in their original DOCX body order."""
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        if hasattr(document, "iter_inner_content"):
            yield from document.iter_inner_content()
            return
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, document)
            elif child.tag == qn("w:tbl"):
                yield Table(child, document)

    @staticmethod
    def _publish(output: Path, source: bytes, markdown: str, assets: dict[str, bytes], report: ImportReport) -> None:
        output.parent.mkdir(parents=True, exist_ok=True)
        temp = Path(tempfile.mkdtemp(prefix=f".{output.name}-", dir=output.parent))
        try:
            (temp / "assets").mkdir()
            (temp / "source.docx").write_bytes(source)
            (temp / "document.md").write_text(markdown, encoding="utf-8", newline="\n")
            for name, content in sorted(assets.items()):
                (temp / "assets" / name).write_bytes(content)
            payload = {
                "dropped": list(report.dropped),
                "normalized": list(report.normalized),
                "preserved": list(report.preserved),
                "source_hash": report.source_hash,
                "unsupported": list(report.unsupported),
            }
            (temp / "import-report.json").write_text(
                json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8", newline="\n"
            )
            backup = output.with_name(output.name + ".previous")
            if backup.exists():
                shutil.rmtree(backup)
            if output.exists():
                os.replace(output, backup)
            os.replace(temp, output)
            if backup.exists():
                shutil.rmtree(backup)
        except Exception:
            if temp.exists():
                shutil.rmtree(temp)
            raise
