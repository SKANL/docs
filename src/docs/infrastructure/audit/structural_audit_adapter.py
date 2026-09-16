from __future__ import annotations

import re
from pathlib import Path
from typing import Any, cast
from zipfile import ZipFile

from docs.domain.review import Issue, ReviewDimension


class StructuralAuditAdapter:
    """Declarative structural checks for DOCX and PDF; prose review is elsewhere."""

    def audit(self, artifact_path: Path, rules: dict[str, object]) -> list[Issue]:
        suffix = artifact_path.suffix.lower()
        rules = self._executable_rules(rules)
        if suffix == ".docx":
            issues = self._audit_docx(artifact_path, rules)
            return self._audit_contract(artifact_path, rules, issues)
        if suffix == ".pdf":
            issues = self._audit_pdf(artifact_path, rules)
            return self._audit_contract(artifact_path, rules, issues)
        return [
            Issue(
                "warning",
                f"No hay auditoría estructural para {suffix or 'este formato'}.",
                "structure.unsupported",
                ReviewDimension.STRUCTURAL,
            )
        ]

    def _audit_docx(self, artifact_path: Path, rules: dict[str, object]) -> list[Issue]:
        from docx import Document

        document = Document(str(artifact_path))
        issues: list[Issue] = []
        headings = [p.text.strip() for p in document.paragraphs if p.style and p.style.name.startswith("Heading") and p.text.strip()]
        expected_headings = self._string_list(rules.get("headings"))
        if expected_headings and not self._ordered(headings, expected_headings):
            issues.append(Issue("error", "Los headings no respetan el orden declarado.", "structure.headings.order", ReviewDimension.STRUCTURAL))
        expected_sections = self._string_list(rules.get("sections"))
        if expected_sections and not self._ordered(headings, expected_sections):
            issues.append(Issue("error", "Las secciones no respetan el orden declarado.", "structure.sections.order", ReviewDimension.STRUCTURAL))

        self._minimum(issues, "tables", len(document.tables), rules.get("tables"), "tabla")
        image_count = len(document.inline_shapes)
        self._minimum(issues, "images", image_count, rules.get("images"), "imagen")
        captions = sum(bool(re.match(r"^(Figura|Figure|Tabla|Table)\s+\d+", p.text.strip(), re.IGNORECASE)) for p in document.paragraphs)
        caption_rules = self._mapping(rules.get("captions"))
        if caption_rules.get("required") and image_count and captions < image_count:
            issues.append(Issue("error", "Faltan captions para imágenes declaradas.", "structure.captions.missing", ReviewDimension.ACCESSIBILITY))
        reference_rules = self._mapping(rules.get("references"))
        if reference_rules.get("required") and not any(re.fullmatch(r"(references|referencias)", heading, re.IGNORECASE) for heading in headings):
            issues.append(Issue("error", "Falta una sección de referencias.", "structure.references.missing", ReviewDimension.STRUCTURAL))
        metadata_rules = self._mapping(rules.get("metadata"))
        for field in self._string_list(metadata_rules.get("required")):
            if not getattr(document.core_properties, field, None):
                issues.append(Issue("error", f"Falta metadato requerido: {field}.", "structure.metadata.missing", ReviewDimension.STRUCTURAL))
        self._page_size(
            issues,
            [(float(s.page_width or 0) / 12700, float(s.page_height or 0) / 12700) for s in document.sections],
            rules,
        )
        return issues

    def _audit_contract(self, artifact_path: Path, rules: dict[str, object], issues: list[Issue]) -> list[Issue]:
        contract: dict[str, Any] = self._mapping(rules.get("template_contract"))
        if not contract:
            contract = cast(dict[str, Any], rules) if any(key in rules for key in (
                "page_geometry", "components", "editable_slots", "required_assets",
                "fidelity_checks", "allowed_degradations",
            )) else {}
        if not contract:
            return issues

        if artifact_path.suffix.lower() == ".docx":
            from docx import Document

            document = Document(str(artifact_path))
            text = "\n".join(p.text for p in document.paragraphs)
            slots = self._docx_slots(artifact_path)
            asset_count = len(document.inline_shapes)
            counts = {
                "headings": sum(bool(p.style and p.style.name.startswith("Heading") and p.text.strip()) for p in document.paragraphs),
                "sections": sum(bool(p.style and p.style.name.startswith("Heading") and p.text.strip()) for p in document.paragraphs),
                "tables": len(document.tables),
                "images": asset_count,
                "captions": sum(bool(re.match(r"^(Figura|Figure|Tabla|Table)\s+\d+", p.text.strip(), re.IGNORECASE)) for p in document.paragraphs),
                "references": sum(bool(re.fullmatch(r"(references|referencias)", p.text.strip(), re.IGNORECASE)) for p in document.paragraphs),
                "metadata": sum(bool(getattr(document.core_properties, field, None)) for field in ("title", "subject", "author", "keywords")),
                "page-count": len(document.sections),
            }
        else:
            text = ""
            slots = set()
            asset_count = 0
            counts = {"page-count": self._pdf_page_count(artifact_path)}

        geometry = self._mapping(contract.get("page_geometry"))
        if geometry:
            self._audit_geometry(issues, artifact_path, geometry)
        for component in contract.get("components", []):
            if isinstance(component, dict):
                kind = str(component.get("kind", "")).casefold()
                if kind in counts:
                    self._audit_component(issues, kind, component, counts[kind])

        for entry in contract.get("required_assets", []):
            if isinstance(entry, dict):
                asset_id = str(entry.get("id", "asset"))
                path_value = entry.get("path") or entry.get("file") or entry.get("source")
                found = bool(path_value and (artifact_path.parent / str(path_value)).is_file())
                if not found and not path_value and str(entry.get("kind", "")).casefold() == "image":
                    found = asset_count > 0
                if not found:
                    self._contract_issue(issues, f"contract.required_assets.{asset_id}",
                                         f"Required asset `{asset_id}` is missing.", contract, asset_id,
                                         required=entry.get("required", True))

        for entry in contract.get("editable_slots", []):
            if isinstance(entry, dict):
                slot_id = str(entry.get("id", "slot"))
                present = slot_id in slots or f"[[slot:{slot_id}]]" in text or f"slot:{slot_id}" in text
                if present:
                    issues.append(Issue("info", f"Editable slot `{slot_id}` is present.",
                                        f"contract.editable_slots.{slot_id}", ReviewDimension.STRUCTURAL))
                else:
                    self._contract_issue(issues, f"contract.editable_slots.{slot_id}",
                                         f"Required editable slot `{slot_id}` is missing.", contract, slot_id,
                                         required=entry.get("required", True))

        for entry in contract.get("fidelity_checks", []):
            if isinstance(entry, dict):
                check_id = str(entry.get("id", "check"))
                actual = counts.get(check_id.casefold())
                if actual is None:
                    self._contract_issue(issues, f"contract.fidelity_checks.{check_id}",
                                         f"Fidelity check `{check_id}` is not executable for this artifact.", contract, check_id,
                                         required=False)
                    continue
                minimum = entry.get("minimum")
                maximum = entry.get("maximum")
                if (isinstance(minimum, (int, float)) and actual < minimum) or (isinstance(maximum, (int, float)) and actual > maximum):
                    self._contract_issue(issues, f"contract.fidelity_checks.{check_id}",
                                         f"Fidelity check `{check_id}` failed: actual value is {actual}.", contract, check_id,
                                         required=entry.get("required", True))
        return issues

    @staticmethod
    def _executable_rules(rules: dict[str, object]) -> dict[str, object]:
        """Project contract components into the legacy rule vocabulary."""
        result = dict(rules)
        contract: dict[str, Any] = StructuralAuditAdapter._mapping(rules.get("template_contract"))
        if not contract and any(key in rules for key in ("page_geometry", "components", "editable_slots", "required_assets", "fidelity_checks")):
            contract = cast(dict[str, Any], rules)
        for component in contract.get("components", []):
            if not isinstance(component, dict):
                continue
            kind = str(component.get("kind", "")).casefold()
            kind = {"heading": "headings", "section": "sections", "table": "tables", "image": "images"}.get(kind, kind)
            if kind and kind not in result:
                if kind in {"headings", "sections"} and isinstance(component.get("items"), list):
                    result[kind] = component["items"]
                else:
                    result[kind] = {key: value for key, value in component.items() if key not in {"id", "kind"}}
        geometry = StructuralAuditAdapter._mapping(contract.get("page_geometry"))
        if geometry and "page_size" not in result:
            size = geometry.get("size")
            sizes = {"a4": (595.28, 841.89), "letter": (612.0, 792.0), "legal": (612.0, 1008.0)}
            if isinstance(size, str) and size.casefold() in sizes:
                result["page_size"] = sizes[size.casefold()]
            elif isinstance(geometry.get("width"), (int, float)) and isinstance(geometry.get("height"), (int, float)):
                result["page_size"] = (geometry["width"], geometry["height"])
        return result

    def _audit_geometry(self, issues: list[Issue], artifact_path: Path, geometry: dict[str, Any]) -> None:
        if artifact_path.suffix.lower() != ".docx":
            return
        from docx import Document

        document = Document(str(artifact_path))
        orientation = geometry.get("orientation")
        margins = geometry.get("margins_cm", geometry.get("margins"))
        for section in document.sections:
            width = float(section.page_width or 0) / 12700
            height = float(section.page_height or 0) / 12700
            if orientation in {"portrait", "landscape"} and ((width > height) != (orientation == "landscape")):
                issues.append(Issue("error", f"Page orientation does not match `{orientation}`.", "contract.page_geometry.orientation", ReviewDimension.STRUCTURAL))
            if isinstance(margins, dict):
                actual = {"top": section.top_margin, "right": section.right_margin, "bottom": section.bottom_margin, "left": section.left_margin}
                if any(key in margins and abs(float(cast(Any, value)) / 360000 - float(margins[key])) > 0.1 for key, value in actual.items()):
                    issues.append(Issue("error", "Page margins do not match the declared geometry.", "contract.page_geometry.margins", ReviewDimension.STRUCTURAL))
                    break

    @staticmethod
    def _audit_component(issues: list[Issue], kind: str, component: dict[str, Any], actual: int) -> None:
        minimum = component.get("minimum", 1 if component.get("required") else 0)
        if isinstance(minimum, int) and actual < minimum:
            if kind in {"headings", "sections"} and component.get("items"):
                code = f"structure.{kind}.order"
            else:
                code = f"structure.{kind}.minimum"
            issues.append(Issue("error", f"Component `{kind}` does not meet the declared minimum.", code, ReviewDimension.STRUCTURAL))

    @staticmethod
    def _contract_issue(issues: list[Issue], code: str, message: str, contract: dict[str, Any], identifier: str, *, required: bool) -> None:
        severity = "error" if required else "warning"
        allowed = " ".join(str(value).casefold() for value in contract.get("allowed_degradations", []))
        if identifier.casefold() in allowed or code.casefold() in allowed:
            severity = "warning"
        issues.append(Issue(severity, message, code, ReviewDimension.STRUCTURAL))

    @staticmethod
    def _docx_slots(path: Path) -> set[str]:
        try:
            with ZipFile(path) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
        except (KeyError, OSError, UnicodeDecodeError):
            return set()
        return set(re.findall(r"(?:w:tag|w:alias)[^>]+w:val=\"([^\"]+)\"", xml))

    @staticmethod
    def _pdf_page_count(path: Path) -> int:
        try:
            import pypdfium2 as pdfium
            document = pdfium.PdfDocument(str(path))
            try:
                return len(document)
            finally:
                document.close()
        except (OSError, RuntimeError, ValueError):
            return 0

    def _audit_pdf(self, artifact_path: Path, rules: dict[str, object]) -> list[Issue]:
        import pypdfium2 as pdfium

        issues: list[Issue] = []
        try:
            document = pdfium.PdfDocument(str(artifact_path))
            try:
                if len(document) == 0:
                    issues.append(Issue("error", "El PDF no contiene páginas.", "structure.pdf.unreadable", ReviewDimension.STRUCTURAL))
                else:
                    self._page_size(issues, [document[index].get_size() for index in range(len(document))], rules)
            finally:
                document.close()
        except (OSError, RuntimeError, ValueError) as exc:
            issues.append(Issue("error", f"PDF no legible: {exc}", "structure.pdf.unreadable", ReviewDimension.STRUCTURAL))
        return issues

    @staticmethod
    def _string_list(value: object) -> list[str]:
        return [str(item) for item in value] if isinstance(value, list) else []

    @staticmethod
    def _mapping(value: object) -> dict[str, Any]:
        return value if isinstance(value, dict) else {}

    @staticmethod
    def _ordered(actual: list[str], expected: list[str]) -> bool:
        position = 0
        for value in actual:
            if position < len(expected) and value.casefold() == expected[position].casefold():
                position += 1
        return position == len(expected)

    def _minimum(self, issues: list[Issue], name: str, actual: int, rule: object, label: str) -> None:
        minimum = self._mapping(rule).get("minimum", 0)
        if isinstance(minimum, int) and actual < minimum:
            issues.append(Issue("error", f"Se requieren al menos {minimum} {label}(s).", f"structure.{name}.minimum", ReviewDimension.STRUCTURAL))

    def _page_size(self, issues: list[Issue], dimensions: list[tuple[float, float]], rules: dict[str, object]) -> None:
        expected = rules.get("page_size")
        if not isinstance(expected, (list, tuple)) or len(expected) != 2:
            return
        for width, height in dimensions:
            if abs(width - expected[0]) > 0.5 or abs(height - expected[1]) > 0.5:
                issues.append(Issue("error", "El tamaño de página no coincide con la regla declarada.", "structure.page_size", ReviewDimension.STRUCTURAL))
                return
